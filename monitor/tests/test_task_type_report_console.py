# -*- coding: utf-8 -*-
"""范围、名单选项、维度键分页及全量 XLSX 的真实服务集成验证。"""

from io import BytesIO

import httpx
from openpyxl import load_workbook
import pytest

from .test_task_type_report_queries import db as db, service_db as service_db
from .test_task_type_report_dimensions import dimension_db as dimension_db
from monitor.app._app import app
from monitor.app.services.cron import task_type_report_export as export_module

URL = "/api/monitor/cron/task-type-report"
DATES = {"start_date": "2026-09-13", "end_date": "2026-09-13"}


async def fetch(params=None, bbk="100", suffix=""):
    headers = {"X-Source-Id": "S"}
    if bbk is not None:
        headers["X-Bbk-Id"] = bbk
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        return await client.get(
            URL + suffix, params=params or DATES, headers=headers
        )


@pytest.fixture
def paged_db(dimension_db):
    for i in range(45):
        user = f"p{i:03}"
        dimension_db.execute(
            "INSERT INTO jkh_user_inf VALUES (?,'2026-09-13','001','01','甲分行','同名支行',?,'L3')",
            (user, f"分页经理{i:03}"),
        )
        dimension_db.execute(
            "INSERT INTO swe_tenant_init_source VALUES (?,'S')", (user,)
        )
        dimension_db.execute(
            "INSERT INTO swe_cron_jobs VALUES (?,?,'S','k1,k2','active',NULL)",
            (user, user),
        )
        dimension_db.execute(
            "INSERT INTO swe_cron_executions VALUES (?,?,?,?,'2026-09-13 10:00:00','success','success',0)",
            (100 + i, user, user, user),
        )
        dimension_db.execute(
            "INSERT INTO swe_cron_subtasks VALUES (?,?)", (user, user)
        )
    return dimension_db


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix", ["", "/options", "/export"])
@pytest.mark.parametrize("bbk", [None, "", " "])
async def test_required_branch_header(service_db, suffix, bbk):
    params = {**DATES, "task_type": "push_plan", "kind": "branches"}
    response = await fetch(params, bbk, suffix)
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "report_validation_error"
    assert service_db.calls == []


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix", ["", "/options", "/export"])
async def test_branch_scope_cannot_be_widened(
    dimension_db, service_db, suffix
):
    response = await fetch(
        {
            **DATES,
            "first_bbk_id": "002",
            "kind": "branches",
            "task_type": "push_plan",
        },
        "001",
        suffix,
    )
    assert response.status_code == 403
    assert response.json()["detail"]["code"] == "report_scope_forbidden"
    assert service_db.calls == []


@pytest.mark.asyncio
async def test_effective_branch_and_cross_scope_user_org(
    dimension_db, service_db
):
    all_rows = await fetch({**DATES, "group_by": "manager"})
    assert all_rows.status_code == 200 and all_rows.json()["total"] == 3
    own = await fetch(
        {**DATES, "group_by": "manager", "task_type": "push_plan"}, "001"
    )
    body = own.json()
    assert body["resolved_filters"]["first_bbk_id"] == "001"
    assert {r["user_id"] for r in body["items"]} == {"alice"}
    dimension_db.execute(
        "UPDATE jkh_user_inf SET org_id='foreign' WHERE user_id='bob'"
    )
    for filters in (
        {"user_id": "bob"},
        {"org_id": "foreign"},
        {"first_bbk_name": "乙分行"},
    ):
        for suffix in ("", "/export"):
            result = await fetch(
                {
                    **DATES,
                    "group_by": "manager",
                    "task_type": "push_plan",
                    **filters,
                },
                "001",
                suffix,
            )
            assert result.status_code == 403, result.text
    narrowed = await fetch({**DATES, "first_bbk_id": "001", "user_id": "bob"})
    assert narrowed.status_code == 403
    named = await fetch(
        {**DATES, "first_bbk_name": "甲分行", "org_id": "foreign"}
    )
    assert named.status_code == 403
    unknown = await fetch({**DATES, "org_id": "no-such-org"}, "001")
    assert unknown.status_code == 200 and unknown.json()["total"] == 0


@pytest.mark.asyncio
async def test_roster_options_and_snapshot(dimension_db, service_db):
    dimension_db.execute(
        "INSERT INTO jkh_user_inf VALUES ('idle','2026-09-13','001','new','','未活动网点',NULL,NULL)"
    )
    dimension_db.execute(
        "INSERT INTO jkh_user_inf VALUES ('blank','2026-09-13','',' ',NULL,NULL,NULL,NULL)"
    )
    args = {"end_date": "2026-09-14", "kind": "branches"}
    all_options = (await fetch(args, suffix="/options")).json()
    assert all_options["sync_date"] == "2026-09-13"
    assert [r["value"] for r in all_options["items"]] == ["001", "002", "003"]
    assert (await fetch(args, "001", "/options")).json()["items"] == [
        {"value": "001", "label": "甲分行"}
    ]
    args["kind"] = "orgs"
    assert (await fetch(args, suffix="/options")).status_code == 422
    orgs = (await fetch(args, "001", "/options")).json()["items"]
    assert orgs == [
        {"value": "01", "label": "同名支行"},
        {"value": "new", "label": "未活动网点"},
    ]
    dimension_db.execute("DELETE FROM jkh_user_inf")
    assert (await fetch(args, "001", "/options")).json() == {
        "sync_date": None,
        "items": [],
    }


@pytest.mark.asyncio
@pytest.mark.parametrize("detail", [False, True])
async def test_pages_match_full_filtered_report(paged_db, service_db, detail):
    params = {
        **DATES,
        "group_by": "manager",
        "keyword": "分页",
        "task_type": "push_plan",
        "skill_detail": str(detail).lower(),
    }
    full = (await fetch(params, "001")).json()
    assert full["total"] == (90 if detail else 45)
    assert full["page"] is None and full["has_more"] is False
    combined = []
    for page in range(1, 7):
        response = await fetch(
            {**params, "page": page, "page_size": 20}, "001"
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["total"] == full["total"]
        assert body["has_more"] is (page * 20 < full["total"])
        assert len(body["items"]) <= 20
        combined.extend(body["items"])
    assert combined == full["items"]
    assert (
        len({(r["user_id"], r["skill_id"]) for r in combined}) == full["total"]
    )
    # Metric SQL contains a page-key restriction, not an arbitrary aggregate LIMIT.
    metric_sql = [
        sql for sql, _ in service_db.calls if "AS suc_execute_job" in sql
    ]
    assert any("user_id IN" in sql for sql in metric_sql)
    if detail:
        assert any(
            "p.user_id = %s AND kd.skill_id = %s" in sql
            for sql in metric_sql
        )


@pytest.mark.asyncio
async def test_keyword_literal_wildcards_and_user_id(dimension_db, service_db):
    dimension_db.execute(
        "UPDATE jkh_user_inf SET user_name='100%_!经理', pst_lvl='特殊岗位' WHERE user_id='alice'"
    )
    base = {**DATES, "group_by": "manager", "task_type": "ask_plan"}
    for keyword in ("100%_!", "特殊岗位", "alice"):
        response = await fetch({**base, "keyword": keyword}, "001")
        assert response.status_code == 200
        assert [r["user_id"] for r in response.json()["items"]] == ["alice"]
    assert (await fetch({**base, "keyword": "' OR 1=1 --"}, "001")).json()[
        "items"
    ] == []
    assert (await fetch({**base, "keyword": "   "}, "001")).json()[
        "total"
    ] == 1
    assert (await fetch({**base, "user_id": "alice"}, "001")).json()[
        "total"
    ] == 1
    assert (await fetch({**base, "user_id": "absent"}, "001")).json()[
        "total"
    ] == 0


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "filters",
    [
        {"page": 1},
        {"page_size": 20},
        {"page": 0, "page_size": 20},
        {"page": 1, "page_size": 101},
        {
            "page": 1,
            "page_size": 20,
            "group_by": "branch",
            "task_type": "push_plan",
        },
        {"page": 1, "page_size": 20, "group_by": "manager"},
        {"keyword": "岗位"},
        {"group_by": "manager", "keyword": "x" * 101},
        {"task_type": "unknown"},
    ],
)
async def test_server_parameter_validation(service_db, filters):
    result = await fetch({**DATES, **filters})
    assert result.status_code == 422
    assert result.json()["detail"]["code"] == "report_validation_error"
    assert service_db.calls == []


@pytest.mark.asyncio
@pytest.mark.parametrize("detail", [False, True])
async def test_full_xlsx_export(paged_db, service_db, detail):
    params = {
        **DATES,
        "group_by": "manager",
        "task_type": "push_plan",
        "keyword": "分页",
        "skill_detail": str(detail).lower(),
    }
    full = (await fetch(params, "001")).json()
    result = await fetch(params, "001", "/export")
    assert result.status_code == 200
    assert result.headers["content-type"] == export_module.XLSX_MEDIA_TYPE
    assert "filename*=UTF-8''" in result.headers["content-disposition"]
    assert (
        result.headers["access-control-expose-headers"]
        == "Content-Disposition"
    )
    workbook = load_workbook(BytesIO(result.content))
    sheet = workbook.worksheets[0]
    assert sheet.max_row == full["total"] + 1
    columns = [column[0] for column in export_module.COLUMNS]
    for source, cells in zip(full["items"], sheet.iter_rows(min_row=2)):
        for field, cell in zip(columns, cells):
            expected = source[field]
            if field.endswith("rate") and expected is not None:
                assert cell.value == pytest.approx(expected / 100)
                assert cell.number_format == "0.00%"
            else:
                assert cell.value == expected
    metadata = list(workbook.worksheets[1].values)
    assert any("不可相加" in str(row) for row in metadata)
    assert any(row == ("keyword", "分页") for row in metadata)
    workbook.close()


@pytest.mark.asyncio
async def test_export_formula_safety_nulls_empty_and_limits(
    dimension_db, service_db, monkeypatch
):
    dimension_db.execute(
        "UPDATE jkh_user_inf SET user_name='=1+1' WHERE user_id='alice'"
    )
    params = {
        **DATES,
        "group_by": "manager",
        "task_type": "push_other",
        "user_id": "alice",
    }
    response = await fetch(params, "001", "/export")
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.worksheets[0]
    assert sheet["F2"].value == "=1+1" and sheet["F2"].data_type == "s"
    columns = [c[0] for c in export_module.COLUMNS]
    assert (
        sheet.cell(2, columns.index("recommended_customers") + 1).value is None
    )
    assert sheet.cell(2, columns.index("suc_execute_job") + 1).value == 1
    assert sheet.cell(2, columns.index("insight_count") + 1).value is None
    assert (
        columns.index("insight_count")
        == columns.index("click_to_insight_rate") + 1
    )
    assert (
        columns.index("phone_count")
        == columns.index("click_to_phone_rate") + 1
    )
    workbook.close()
    for override in ({"task_type": None}, {"page": 1, "page_size": 20}):
        result = await fetch({**params, **override}, "001", "/export")
        assert result.status_code == 422
    empty = await fetch({**params, "user_id": "absent"}, "001", "/export")
    workbook = load_workbook(BytesIO(empty.content))
    assert workbook.worksheets[0].max_row == 1
    workbook.close()
    monkeypatch.setattr(export_module, "MAX_EXPORT_ROWS", 0)
    limited = await fetch(params, "001", "/export")
    assert limited.status_code == 413
    assert limited.json()["detail"]["code"] == "report_export_too_large"


@pytest.mark.asyncio
async def test_page_boundary_splits_one_manager_skills(paged_db, service_db):
    params = {
        **DATES,
        "group_by": "manager",
        "skill_detail": "true",
        "task_type": "ask_plan",
        "user_id": "p000",
        "page_size": 1,
    }
    first = (await fetch({**params, "page": 1}, "001")).json()
    second = (await fetch({**params, "page": 2}, "001")).json()
    assert first["total"] == second["total"] == 2
    assert first["has_more"] is True and second["has_more"] is False
    assert (
        first["items"][0]["user_id"] == second["items"][0]["user_id"] == "p000"
    )
    assert first["items"][0]["skill_id"] == "k1"
    assert second["items"][0]["skill_id"] == "k2"
    assert all(
        item["suc_execute_job"] == item["read_tasks"] == 0
        and item["active_manager_count"] is None
        for item in first["items"] + second["items"]
    )
