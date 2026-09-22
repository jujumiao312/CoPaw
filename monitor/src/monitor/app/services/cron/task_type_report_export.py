# -*- coding: utf-8 -*-
"""沿用项目 openpyxl 工具生成真实 XLSX，全量数据由报表服务提供。

列集合与当前导出的报表一一对应：维度列跟随 ``group_by``，技能明细只追加技能
名称，任务类型天然不产出的指标（push_other 的客户方案、ask_plan 的活跃人数）
不导出，避免整列恒空的字段混进文件。
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from .task_type_report import ReportError
from .task_type_report_rows import (
    LABELS as TASK_TYPE_LABELS,
    NULL_FIELDS,
    RATIOS,
)

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
MAX_EXPORT_ROWS = 50000
# 每种维度报表实际展示的标识列，overall 不带任何维度列。
DIMENSION_COLUMNS = {
    "overall": (),
    "branch": ("first_bbk_id", "first_bbk_name"),
    "org": ("first_bbk_id", "first_bbk_name", "org_id", "org_name"),
    "manager": (
        "first_bbk_id",
        "first_bbk_name",
        "org_id",
        "org_name",
        "user_id",
        "user_name",
        "sapid",
        "pst_lvl",
    ),
}
# 技能明细只导出技能名称；技能 ID 由 Console 作为名称附注展示，不单列。
SKILL_COLUMNS = ("cn_name",)
MANAGER_TASK_COLUMNS = (
    ("active_task_count", "当前活跃任务数"),
    ("paused_task_count", "当前暂停任务数"),
)
# 指标列顺序固定；任务类型天然不产出的指标在组列时剔除。
METRIC_COLUMNS = (
    ("skill_count", "技能数"),
    ("permission_manager_count", "有权限客户经理数"),
    ("active_manager_count", "活跃客户经理数"),
    ("suc_execute_job", "成功任务数"),
    ("read_tasks", "已查看任务数"),
    ("read_rate", "任务查看率（%）"),
    ("recommended_customers", "方案客户数"),
    ("read_customer_count", "已查看方案客户数"),
    ("plan_read_rate", "方案查看率（%）"),
    ("insight_customer_count", "洞察客户数"),
    ("click_to_insight_rate", "洞察覆盖率（%）"),
    ("insight_count", "点击客户洞察总次数"),
    ("phone_customer_count", "电访客户数"),
    ("click_to_phone_rate", "电访覆盖率（%）"),
    ("phone_count", "点击去电访总次数"),
)
LABELS = {
    "first_bbk_id": "分行号",
    "first_bbk_name": "分行名称",
    "org_id": "网点号",
    "org_name": "网点名称",
    "user_id": "客户经理ID",
    "user_name": "客户经理姓名",
    "sapid": "SAP号",
    "pst_lvl": "SAP岗位",
    "cn_name": "技能名称",
    "task_type_name": "任务类型",
    **dict(METRIC_COLUMNS),
    **dict(MANAGER_TASK_COLUMNS),
}


def missing_metrics(task_type: str) -> set[str]:
    """该任务类型天然缺失的指标：置空字段，以及分子或分母缺失的比例。"""
    missing = set(NULL_FIELDS.get(task_type, ()))
    for field, (numerator, denominator) in RATIOS.items():
        if numerator in missing or denominator in missing:
            missing.add(field)
    return missing


def column_labels(
    group_by: str, skill_detail: bool, task_type: str, *, snapshot: bool = False
) -> list[tuple[str, str]]:
    """按对应报表取列，只保留该维度与该任务类型真实产出的字段。"""
    fields = list(DIMENSION_COLUMNS[group_by])
    if skill_detail:
        fields.extend(SKILL_COLUMNS)
    fields.append("task_type_name")
    missing = missing_metrics(task_type)
    metrics = list(METRIC_COLUMNS)
    if snapshot:
        if skill_detail:
            missing.update(("skill_count", "permission_manager_count"))
        if group_by == "manager":
            metrics[1:3] = MANAGER_TASK_COLUMNS
    fields.extend(field for field, _ in metrics if field not in missing)
    return [(field, LABELS[field]) for field in fields]


def _cell(sheet, value, percentage=False):
    cell = WriteOnlyCell(sheet, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    elif value is not None:
        if percentage:
            cell.value = value / 100
            cell.number_format = "0.00%"
        else:
            cell.number_format = "#,##0"
    return cell


def _row_value(row, field):
    """任务类型名称以码值映射为准，兼容落盘表短名或旧值。"""
    if field != "task_type_name":
        return getattr(row, field)
    return TASK_TYPE_LABELS.get(getattr(row, "task_type"), getattr(row, field))


def export_task_type_report(report, task_type: str) -> bytes:
    if len(report.items) > MAX_EXPORT_ROWS:
        raise ReportError(
            413,
            "report_export_too_large",
            "导出超过50000行，请缩小时间或机构范围。",
        )
    columns = column_labels(
        report.group_by,
        report.skill_detail,
        task_type,
        snapshot=report.consistency == "snapshot",
    )
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(
        "技能明细" if report.skill_detail else "统计报表"
    )
    sheet.freeze_panes = "A2"
    headers = [_cell(sheet, label) for _, label in columns]
    for cell in headers:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    sheet.append(headers)
    for row in report.items:
        sheet.append(
            [
                _cell(sheet, _row_value(row, field), field.endswith("rate"))
                for field, _ in columns
            ]
        )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
