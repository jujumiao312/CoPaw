# -*- coding: utf-8 -*-
"""落盘快照接口：SQLite 验证过滤/排序/分页/装配，ASGI 验证契约与导出。

SQLite 只用于验证逻辑与 SQL 形状，不代表 TDSQL 方言与性能验证。
"""

import sqlite3
from datetime import date
from io import BytesIO
import re
from types import SimpleNamespace

import httpx
import pytest
from openpyxl import load_workbook

from monitor.app._app import app
from monitor.app.database import schema
from monitor.app.models.task_type_report import (
    COMMON_METRIC_FIELDS,
    COMMON_METRIC_INT_FIELDS,
    ReportOptionsParams,
)
from monitor.app.models.task_type_snapshot import SnapshotReportParams
from monitor.app.services.cron import task_type_report as cron_report
from monitor.app.services.report.task_type_snapshot import (
    SNAPSHOT_SOURCE_COLUMNS,
    TaskTypeSnapshotService,
)

BASE_URL = "/api/monitor/report/task-type"
HEADERS = {"X-Source-Id": "S", "X-Bbk-Id": "100"}
DATE = "2026-09-17"
DATE_VALUE = date(2026, 9, 17)
TASK_TYPES = ("push_plan", "ask_plan", "push_other")
LABELS = {
    "push_plan": "推送(名单+方案)",
    "ask_plan": "主动提问(名单+方案)",
    "push_other": "推送(非名单方案)",
}


class SqliteConnection:
    """把同步 sqlite3 包成报表服务需要的最小异步连接接口。"""

    def __init__(self, connection: sqlite3.Connection):
        self._connection = connection
        self._connection.row_factory = sqlite3.Row

    @property
    def is_connected(self) -> bool:
        return True

    async def fetch_all(self, sql: str, params=()):
        cursor = self._connection.execute(
            sql.replace("%s", "?"), tuple(params or ())
        )
        return [dict(row) for row in cursor.fetchall()]

    async def fetch_one(self, sql: str, params=()):
        rows = await self.fetch_all(sql, params)
        return rows[0] if rows else None

    async def execute(self, sql: str, params=()):
        self._connection.execute(sql.replace("%s", "?"), tuple(params or ()))
        self._connection.commit()


COMMON_SCHEMA_COLUMNS = ",\n    ".join(
    f"{field} {'INTEGER' if field in COMMON_METRIC_INT_FIELDS else 'REAL'}"
    for field in COMMON_METRIC_FIELDS
)
SCHEMA = f"""
CREATE TABLE swe_rm_claw_list_ind_stat(
    dw_dat_dt TEXT, source_id TEXT,
    rpt_combo TEXT, job_type TEXT, frs_bbk_org_id TEXT DEFAULT '',
    brn_org_id TEXT DEFAULT '', cm_id TEXT DEFAULT '',
    skill_id TEXT DEFAULT '', frs_bbk_org_nm TEXT DEFAULT '',
    brn_org_nm TEXT DEFAULT '', cm_nm TEXT DEFAULT '',
    pst_lvl TEXT DEFAULT '', skill_nm TEXT DEFAULT '',
    job_type_nm TEXT DEFAULT '', skill_cnt INTEGER DEFAULT 0,
    active_manager_cnt INTEGER, active_job_cnt INTEGER, paused_job_cnt INTEGER,
    suc_execute_job INTEGER DEFAULT 0,
    read_tasks INTEGER DEFAULT 0, read_rate REAL,
    recommended_customers INTEGER, read_customer_cnt INTEGER,
    plan_read_rate REAL, insight_customer_cnt INTEGER,
    click_to_insight_rate REAL, insight_cnt INTEGER,
    phone_customer_cnt INTEGER, click_to_phone_rate REAL,
    phone_cnt INTEGER,
    {COMMON_SCHEMA_COLUMNS},
    stat_start_dt TEXT, stat_end_dt TEXT);
CREATE TABLE jkh_user_inf(user_id TEXT, sync_date TEXT, first_bbk_id TEXT,
    org_id TEXT, first_bbk_nm TEXT, org_nm TEXT, user_name TEXT,
    pst_lvl TEXT);
CREATE TABLE swe_tenant_init_source(tenant_id TEXT, source_id TEXT);
"""

INSERT_SQL = (
    "INSERT INTO swe_rm_claw_list_ind_stat ("
    "dw_dat_dt, source_id, rpt_combo, job_type, frs_bbk_org_id, "
    "brn_org_id, cm_id, skill_id, frs_bbk_org_nm, brn_org_nm, cm_nm, "
    "pst_lvl, skill_nm, job_type_nm, skill_cnt, active_manager_cnt, "
    "suc_execute_job, read_tasks, read_rate, recommended_customers, "
    "stat_start_dt, stat_end_dt) VALUES ("
    ":prt_dt, :source_id, :rpt_combo, :task_type, :first_bbk_id, "
    ":org_id, :user_id, :skill_id, :first_bbk_nm, :org_nm, :user_name, "
    ":pst_lvl, :cn_name, :task_type_name, :skill_cnt, "
    ":active_manager_cnt, :suc_execute_job, :read_tasks, :read_rate, "
    ":recommended_customers, :stat_start_dt, :stat_end_dt)"
)


def snapshot_row(combo, dims, task_type, **metrics):
    """构造一行落盘快照；不适用维度的空串由 dims 缺省值提供。"""
    values = {
        "prt_dt": DATE,
        "source_id": "S",
        "rpt_combo": combo,
        "task_type": task_type,
        "task_type_name": LABELS[task_type],
        "first_bbk_id": "",
        "org_id": "",
        "user_id": "",
        "skill_id": "",
        "first_bbk_nm": (
            "甲分行" if dims.get("first_bbk_id") == "001" else "乙分行"
        ),
        "org_nm": "同名支行",
        "user_name": "alice" if dims.get("user_id") == "alice" else "bob",
        "pst_lvl": "L2" if dims.get("user_id") == "alice" else "L1",
        "cn_name": "技能一" if dims.get("skill_id") else "",
        "stat_start_dt": "2026-09-01",
        "stat_end_dt": DATE,
        "skill_cnt": 1,
        "suc_execute_job": 3,
        "read_tasks": 2,
        "read_rate": 66.67,
        "active_manager_cnt": None,
        "recommended_customers": None,
    }
    values.update(dims)
    values.update(metrics)
    return values


def seed(connection: sqlite3.Connection) -> int:
    """只建快照、名单、权限来源三张表，准备六个组合数据。"""
    connection.executescript(SCHEMA)
    connection.executemany(
        "INSERT INTO jkh_user_inf VALUES(?,?,?,?,?,?,?,?)",
        [
            ("alice", DATE, "001", "01", "甲分行", "同名支行", "alice", "L2"),
            ("bob", DATE, "002", "01", "乙分行", "同名支行", "bob", "L1"),
            ("carol", DATE, "003", "03", "丙分行", "空支行", "carol", "L3"),
        ],
    )
    connection.executemany(
        "INSERT INTO swe_tenant_init_source VALUES(?,?)",
        [("alice", "S"), ("alice", "S"), ("bob", "S"), ("bob", "OTHER")],
    )
    rows = []
    for task_type in TASK_TYPES:
        rows.append(snapshot_row("overall", {}, task_type))
        for bbk, user in (("001", "alice"), ("002", "bob")):
            rows.append(
                snapshot_row("branch", {"first_bbk_id": bbk}, task_type)
            )
            rows.append(
                snapshot_row(
                    "org", {"first_bbk_id": bbk, "org_id": "01"}, task_type
                )
            )
            rows.append(
                snapshot_row(
                    "manager",
                    {"first_bbk_id": bbk, "org_id": "01", "user_id": user},
                    task_type,
                    active_manager_cnt=None if task_type == "ask_plan" else 1,
                    recommended_customers=(
                        None if task_type == "push_other" else 4
                    ),
                )
            )
            rows.append(
                snapshot_row(
                    "branch_skill",
                    {"first_bbk_id": bbk, "skill_id": "k1"},
                    task_type,
                )
            )
            rows.append(
                snapshot_row(
                    "manager_skill",
                    {
                        "first_bbk_id": bbk,
                        "org_id": "01",
                        "user_id": user,
                        "skill_id": "k1",
                    },
                    task_type,
                )
            )
    connection.executemany(INSERT_SQL, rows)
    common_assignments = ", ".join(
        f"{field} = ?" for field in COMMON_METRIC_FIELDS
    )
    common_values = tuple(
        index
        if field in COMMON_METRIC_INT_FIELDS
        else (index + 0.25) / 100
        if field == "vld_ctc_cust_rate"
        else index + 0.25
        for index, field in enumerate(COMMON_METRIC_FIELDS, 1)
    )
    connection.execute(
        f"UPDATE swe_rm_claw_list_ind_stat SET {common_assignments}",
        common_values,
    )
    connection.execute(
        "INSERT INTO swe_rm_claw_list_ind_stat ("
        "dw_dat_dt, source_id, rpt_combo, job_type, frs_bbk_org_id, "
        "brn_org_id, cm_id, skill_id, stat_start_dt, stat_end_dt) VALUES "
        "(?, 'OTHER', 'overall', 'push_plan', '', '', '', '', ?, ?)",
        (DATE, "2026-09-01", DATE),
    )
    connection.commit()
    return len(rows)


@pytest.fixture
def env(monkeypatch):
    raw = sqlite3.connect(":memory:")
    row_count = seed(raw)
    connection = SqliteConnection(raw)
    monkeypatch.setattr(cron_report, "get_db_connection", lambda: connection)
    yield SimpleNamespace(raw=raw, connection=connection, row_count=row_count)
    raw.close()


def params(**overrides) -> SnapshotReportParams:
    values = {"end_date": DATE}
    values.update(overrides)
    return SnapshotReportParams(**values)


def options(kind: str, bbk: str | None = None) -> ReportOptionsParams:
    return ReportOptionsParams(
        end_date=DATE_VALUE, kind=kind, first_bbk_id=bbk
    )


def service() -> TaskTypeSnapshotService:
    return TaskTypeSnapshotService()


async def request(path=BASE_URL, query=None, headers=None):
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        return await client.get(
            path, params=query or {}, headers=headers or HEADERS
        )


@pytest.mark.asyncio
async def test_combo_mapping_and_dimension_restore(env):
    report = await service().get_report(
        params(group_by="manager", skill_detail=True), "S", "100"
    )
    assert report.rpt_combo == "manager_skill"
    assert report.consistency == "snapshot"
    assert report.metric_version == "jkh_task_report_v1_snapshot"
    assert report.sync_date == DATE
    assert report.start_date.isoformat() == "2026-09-01"
    assert report.total == 6
    row = report.items[0]
    assert row.first_bbk_id == "001"
    assert row.org_id == "01"
    assert row.user_id == "alice"
    assert row.sapid == "alice"
    assert row.skill_id == "k1"
    assert row.cn_name == "技能一"
    assert row.permission_manager_count is None
    assert row.skill_count is None
    assert row.task_type == "push_plan"
    assert row.task_type_name == "推送(名单+方案)"
    assert "skill_rows_not_additive" in report.warnings
    assert "snapshot_month_to_date" in report.warnings


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "group_by,skill_detail",
    [
        ("branch", False),
        ("org", False),
        ("manager", False),
        ("branch", True),
        ("manager", True),
    ],
)
async def test_common_metrics_return_for_each_available_dimension(
    env, group_by, skill_detail
):
    report = await service().get_report(
        params(group_by=group_by, skill_detail=skill_detail), "S", "100"
    )
    assert report.items
    for row in report.items:
        assert all(
            getattr(row, field) is not None
            for field in COMMON_METRIC_FIELDS
        )
        assert row.vld_ctc_cust_qty == 1
        assert row.vld_ctc_cust_rate == pytest.approx(2.25)
        assert row.wlth_prod_buy_cm_qty_t1 == 3
        assert row.fnd_inc_t14 == 46.25


@pytest.mark.asyncio
async def test_overall_rows_have_no_dimensions(env):
    report = await service().get_report(params(), "S", "100")
    assert report.rpt_combo == "overall"
    assert report.total == 3
    for row in report.items:
        assert row.first_bbk_id is None
        assert row.org_id is None
        assert row.user_id is None
        assert row.skill_id is None
        assert row.first_bbk_name is None
    assert report.batch.row_total == env.row_count


@pytest.mark.asyncio
async def test_task_type_filter_and_stable_order(env):
    single = await service().get_report(
        params(group_by="branch", task_type="ask_plan"), "S", "100"
    )
    assert [row.task_type for row in single.items] == ["ask_plan"] * 2
    full = await service().get_report(params(group_by="branch"), "S", "100")
    assert [row.task_type for row in full.items] == [
        "push_plan",
        "ask_plan",
        "push_other",
    ] * 2
    assert [row.first_bbk_id for row in full.items[:3]] == ["001"] * 3
    assert full.items[0].permission_manager_count == 1
    assert full.items[3].permission_manager_count == 1


@pytest.mark.asyncio
async def test_null_metrics_stay_null(env):
    report = await service().get_report(params(group_by="manager"), "S", "100")
    rows = {(row.task_type, row.user_id): row for row in report.items}
    assert rows[("ask_plan", "alice")].active_manager_count is None
    assert rows[("push_other", "alice")].recommended_customers is None
    assert rows[("push_plan", "alice")].recommended_customers == 4
    assert rows[("push_plan", "alice")].active_manager_count is None


@pytest.mark.asyncio
async def test_pagination_returns_total_and_has_more(env):
    report = await service().get_report(
        params(
            group_by="manager",
            task_type="push_plan",
            page=2,
            page_size=1,
        ),
        "S",
        "100",
    )
    assert report.total == 2
    assert report.page == 2 and report.page_size == 1
    assert report.has_more is False
    assert [row.user_id for row in report.items] == ["bob"]
    first_page = await service().get_report(
        params(
            group_by="manager",
            task_type="push_plan",
            page=1,
            page_size=1,
        ),
        "S",
        "100",
    )
    assert [row.user_id for row in first_page.items] == ["alice"]
    assert first_page.has_more is True


@pytest.mark.asyncio
async def test_keyword_matches_name_id_and_level(env):
    for keyword, expected in (
        ("ali", "alice"),
        ("L1", "bob"),
        ("bob", "bob"),
        ("1_", ""),
    ):
        report = await service().get_report(
            params(group_by="manager", task_type="push_plan", keyword=keyword),
            "S",
            "100",
        )
        assert [row.user_id for row in report.items] == (
            [expected] if expected else []
        )


@pytest.mark.asyncio
async def test_filter_scope_must_match_combo(env):
    with pytest.raises(cron_report.ReportError) as excinfo:
        await service().get_report(params(first_bbk_id="001"), "S", "100")
    assert excinfo.value.code == "report_filter_not_supported"
    assert excinfo.value.status_code == 422
    with pytest.raises(cron_report.ReportError) as excinfo:
        await service().get_report(
            params(group_by="branch", org_id="01"), "S", "100"
        )
    assert excinfo.value.code == "report_filter_not_supported"


@pytest.mark.asyncio
async def test_scope_forbidden_and_ambiguous_name(env):
    with pytest.raises(cron_report.ReportError) as excinfo:
        await service().get_report(
            params(group_by="branch", first_bbk_id="003"), "S", "001"
        )
    assert excinfo.value.status_code == 403
    with pytest.raises(cron_report.ReportError) as excinfo:
        await service().get_report(
            params(group_by="org", org_name="同名支行"), "S", "100"
        )
    assert excinfo.value.code == "organization_name_ambiguous"


@pytest.mark.asyncio
async def test_name_resolution_narrows_scope(env):
    report = await service().get_report(
        params(group_by="branch", first_bbk_name="甲分行"), "S", "100"
    )
    assert report.resolved_filters.first_bbk_id == "001"
    assert {row.first_bbk_id for row in report.items} == {"001"}
    empty = await service().get_report(
        params(group_by="branch", first_bbk_name="不存在分行"), "S", "100"
    )
    assert empty.items == [] and empty.total == 0
    assert "no_matching_organization" in empty.warnings


@pytest.mark.asyncio
async def test_missing_and_loading_batch(env):
    with pytest.raises(cron_report.ReportError) as excinfo:
        await service().get_report(params(end_date="2026-09-30"), "S", "100")
    assert excinfo.value.status_code == 404
    assert excinfo.value.code == "report_snapshot_not_found"
    with pytest.raises(cron_report.ReportError) as excinfo:
        await service().get_report(params(end_date="2026-09-18"), "S", "100")
    assert excinfo.value.status_code == 404
    assert excinfo.value.code == "report_snapshot_not_found"


@pytest.mark.asyncio
async def test_source_isolation(env):
    report = await service().get_report(params(), "OTHER", "100")
    assert report.total == 1
    assert report.batch.source_id == "OTHER"


@pytest.mark.asyncio
async def test_report_without_batch_table_keeps_permission_counts(env):
    tables = {
        row[0]
        for row in env.raw.execute(
            "SELECT name FROM sqlite_master WHERE type = 'table'"
        )
    }
    assert "swe_task_type_report_batch" not in tables
    response = await request(
        query={
            "end_date": DATE,
            "group_by": "org",
            "task_type": "push_plan",
        }
    )
    assert response.status_code == 200
    body = response.json()
    assert body["sync_date"] == DATE
    assert body["batch"]["sync_date"] == DATE
    assert body["batch"]["loaded_at"] is None
    assert [row["permission_manager_count"] for row in body["items"]] == [1, 1]


@pytest.mark.asyncio
async def test_dates_and_status(env):
    dates = await service().get_dates("S")
    assert [item.prt_dt.isoformat() for item in dates.items] == [
        DATE,
    ]
    assert dates.latest_ready_prt_dt.isoformat() == DATE
    scoped = await service().get_dates("S", "2026-09")
    assert len(scoped.items) == 1
    status = await service().get_status(params(), "S", "100")
    assert status.batch.status == "ready"
    counts = {item.rpt_combo: item.row_cnt for item in status.combo_counts}
    assert counts["manager"] == 6
    assert counts["overall"] == 3
    assert status.missing_combos == ["org_skill"]


@pytest.mark.asyncio
async def test_options_follow_batch_snapshot(env):
    branches = await service().get_options(options("branches"), "S", "100")
    assert [item.value for item in branches.items] == ["001", "002", "003"]
    assert branches.sync_date == DATE
    orgs = await service().get_options(options("orgs", "001"), "S", "100")
    assert [item.value for item in orgs.items] == ["01"]


@pytest.mark.asyncio
async def test_http_contract_and_validation(env):
    response = await request(query={"end_date": DATE, "group_by": "branch"})
    assert response.status_code == 200
    body = response.json()
    assert body["metric_version"] == "jkh_task_report_v1_snapshot"
    assert body["consistency"] == "snapshot"
    assert body["ratio_unit"] == "percent"
    assert body["read_evidence"]["ask"] == "assumed_from_success"
    assert body["resolved_filters"] == {
        "first_bbk_id": None,
        "org_id": None,
    }
    assert body["rpt_combo"] == "branch"
    assert len(body["items"]) == 6
    default_start = await request(query={"end_date": DATE})
    assert default_start.status_code == 200
    assert default_start.json()["start_date"] == "2026-09-01"
    bad_start = await request(
        query={"end_date": DATE, "start_date": "2026-09-02"}
    )
    assert bad_start.status_code == 422
    assert bad_start.json()["detail"]["code"] == "report_validation_error"
    missing_header = await request(headers={"X-Source-Id": "S"})
    assert missing_header.status_code == 422
    not_found = await request(query={"end_date": "2026-09-30"})
    assert not_found.status_code == 404
    not_ready = await request(query={"end_date": "2026-09-18"})
    assert not_ready.status_code == 404
    assert not_ready.json()["detail"]["code"] == "report_snapshot_not_found"
    filter_scope = await request(
        query={"end_date": DATE, "first_bbk_id": "001"}
    )
    assert filter_scope.status_code == 422
    assert (
        filter_scope.json()["detail"]["code"] == "report_filter_not_supported"
    )


@pytest.mark.asyncio
async def test_http_export_xlsx(env):
    response = await request(
        path=f"{BASE_URL}/export",
        query={
            "end_date": DATE,
            "group_by": "branch",
            "task_type": "push_plan",
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.title == "统计报表"
    assert sheet.max_row == 3
    assert sheet["A1"].value == "分行号"
    assert sheet["A2"].value == "001"
    # 有权限客户经理数不在数据源里，导出必须带上接口实时算出来的值
    headers = [cell.value for cell in sheet[1]]
    assert "有权限客户经理数" in headers
    permission_column = headers.index("有权限客户经理数") + 1
    assert headers[-46] == "强接触客户数"
    assert headers[-1] == "基金中收(T+14)"
    assert sheet.cell(row=2, column=len(headers) - 45).value == 1
    rate_column = headers.index("强接触客户率") + 1
    rate_cell = sheet.cell(row=2, column=rate_column)
    assert rate_cell.value == pytest.approx(0.0225)
    assert rate_cell.number_format == "0.00%"
    assert sheet.cell(row=2, column=len(headers)).value == 46.25
    assert sheet.cell(row=2, column=permission_column).value == 1
    skill_export = await request(
        path=f"{BASE_URL}/export",
        query={
            "end_date": DATE,
            "group_by": "branch",
            "skill_detail": True,
            "task_type": "push_plan",
        },
    )
    skill_sheet = load_workbook(BytesIO(skill_export.content)).active
    assert skill_sheet.title == "技能明细"
    assert skill_sheet["C1"].value == "技能名称"
    without_type = await request(
        path=f"{BASE_URL}/export", query={"end_date": DATE}
    )
    assert without_type.status_code == 422
    with_page = await request(
        path=f"{BASE_URL}/export",
        query={
            "end_date": DATE,
            "group_by": "manager",
            "task_type": "push_plan",
            "page": 1,
            "page_size": 10,
        },
    )
    assert with_page.status_code == 422


@pytest.mark.asyncio
@pytest.mark.parametrize("group", ["overall", "branch", "org", "manager"])
async def test_export_task_type_names_are_canonical_for_every_group(
    env, group
):
    env.raw.execute(
        "UPDATE swe_rm_claw_list_ind_stat SET job_type_nm = job_type"
    )
    for task_type in TASK_TYPES:
        exported = await request(
            path=f"{BASE_URL}/export",
            query={
                "end_date": DATE,
                "group_by": group,
                "task_type": task_type,
            },
        )
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        column = headers.index("任务类型") + 1
        values = [
            sheet.cell(row=row, column=column).value
            for row in range(2, sheet.max_row + 1)
        ]
        assert values and set(values) == {LABELS[task_type]}
        workbook.close()


@pytest.mark.asyncio
async def test_http_dates_status_and_options(env):
    dates = await request(path=f"{BASE_URL}/dates")
    assert dates.status_code == 200
    assert dates.json()["latest_ready_prt_dt"] == DATE
    scoped = await request(
        path=f"{BASE_URL}/dates", query={"month": "2026-09", "limit": 1}
    )
    assert len(scoped.json()["items"]) == 1
    status = await request(path=f"{BASE_URL}/status", query={"end_date": DATE})
    assert status.status_code == 200
    assert status.json()["missing_combos"] == ["org_skill"]
    options_response = await request(
        path=f"{BASE_URL}/options",
        query={"end_date": DATE, "kind": "branches"},
    )
    assert options_response.status_code == 200
    assert [item["value"] for item in options_response.json()["items"]] == [
        "001",
        "002",
        "003",
    ]
    bad_month = await request(
        path=f"{BASE_URL}/dates", query={"month": "2026-13"}
    )
    assert bad_month.status_code == 422
    bad_kind = await request(
        path=f"{BASE_URL}/options",
        query={"end_date": DATE, "kind": "unknown"},
    )
    assert bad_kind.status_code == 422


def seed_high_source(connection: sqlite3.Connection) -> None:
    """造一批按高斯落数口径写入的行：不适用维度写 'ALL'。"""
    connection.executemany(
        "INSERT INTO swe_rm_claw_list_ind_stat ("
        "dw_dat_dt, source_id, rpt_combo, job_type, frs_bbk_org_id,"
        " brn_org_id, cm_id, skill_id, frs_bbk_org_nm, brn_org_nm,"
        " cm_nm, pst_lvl, skill_nm, job_type_nm, active_manager_cnt,"
        " skill_cnt) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)",
        [
            (
                DATE, "G", "overall", "push_plan", "ALL", "ALL", "ALL", "ALL",
                "ALL", "ALL", "ALL", "ALL", "ALL", LABELS["push_plan"], 7,
            ),
            (
                DATE, "G", "branch", "push_plan", "001", "ALL", "ALL", "ALL",
                "甲分行", "ALL", "ALL", "ALL", "ALL", LABELS["push_plan"], 7,
            ),
            (
                DATE, "G", "manager", "push_plan", "001", "01", "alice", "ALL",
                "甲分行", "同名支行", "alice", "L2", "ALL",
                LABELS["push_plan"], None,
            ),
            (
                DATE, "G", "branch", "push_plan", "ALL", "ALL", "ALL", "ALL",
                "ALL", "ALL", "ALL", "ALL", "ALL", LABELS["push_plan"], 7,
            ),
        ],
    )
    connection.commit()


@pytest.mark.asyncio
async def test_not_applicable_all_never_reaches_client(env):
    """高斯写 'ALL' 表示“本组合用不到的维度”，出参只能是 null。"""
    seed_high_source(env.raw)
    overall = await service().get_report(params(), "G", "100")
    for row in overall.items:
        assert row.first_bbk_id is None
        assert row.org_id is None
        assert row.user_id is None
        assert row.skill_id is None
        assert row.first_bbk_name is None
        assert row.cn_name is None
        assert row.pst_lvl is None
    branch = await service().get_report(params(group_by="branch"), "G", "100")
    by_bbk = {row.first_bbk_id: row for row in branch.items}
    assert by_bbk["001"].first_bbk_name == "甲分行"
    assert by_bbk["001"].org_id is None
    assert by_bbk["001"].skill_id is None
    # 组合维度列自己也是 'ALL' 时（上游异常数据）同样不能当机构号返回
    assert None in by_bbk
    assert by_bbk[None].first_bbk_name is None
    manager = await service().get_report(params(group_by="manager"), "G", "100")
    assert manager.items[0].user_id == "alice"
    assert manager.items[0].sapid == "alice"
    assert manager.items[0].pst_lvl == "L2"
    assert manager.items[0].cn_name is None
    response = await request(
        query={"end_date": DATE, "group_by": "manager"},
        headers={"X-Source-Id": "G", "X-Bbk-Id": "100"},
    )
    assert response.status_code == 200
    assert "ALL" not in response.text


def test_selected_columns_exist_in_table_ddl():
    """服务查询的列必须都建在表里：SQLite 夹具发现不了列名写错的情况。"""

    def declared_columns(ddl: str) -> set[str]:
        return set(
            re.findall(
                r"^\s*`?([a-z0-9_]+)`?\s+"
                r"(?:BIGINT|DATE|DATETIME|VARCHAR|DECIMAL|CHAR)\b",
                ddl,
                re.M,
            )
        )

    snapshot_columns = declared_columns(
        schema.CREATE_TASK_TYPE_REPORT_SNAPSHOT_TABLE
    )
    assert set(SNAPSHOT_SOURCE_COLUMNS) <= snapshot_columns


@pytest.mark.asyncio
async def test_dates_aggregate_rows_and_filter_source_month_limit(env):
    env.raw.executemany(
        "INSERT INTO swe_rm_claw_list_ind_stat "
        "(dw_dat_dt, source_id, rpt_combo, job_type) "
        "VALUES (?, ?, 'overall', 'push_plan')",
        [
            ("2026-09-20", "S"),
            ("2026-09-20", "S"),
            ("2026-10-01", "S"),
            ("2026-09-21", "OTHER"),
        ],
    )
    dates = await service().get_dates("S", "2026-09", limit=1)
    assert dates.latest_ready_prt_dt.isoformat() == "2026-09-20"
    assert len(dates.items) == 1
    assert dates.items[0].row_total == 2
    assert dates.items[0].status == "ready"
    assert dates.items[0].loaded_at is None
    empty = await service().get_dates("ABSENT")
    assert empty.items == [] and empty.latest_ready_prt_dt is None


@pytest.mark.asyncio
async def test_permission_counts_use_run_date_and_source(env):
    env.raw.execute(
        "UPDATE jkh_user_inf SET sync_date = '2026-09-16' WHERE user_id = 'alice'"
    )
    env.raw.execute(
        "INSERT INTO swe_tenant_init_source VALUES ('carol', 'OTHER')"
    )
    report = await service().get_report(params(), "S", "100")
    assert all(row.permission_manager_count == 1 for row in report.items)
    other = await service().get_report(params(), "OTHER", "100")
    assert other.items[0].permission_manager_count == 2


@pytest.mark.asyncio
async def test_options_prefer_requested_roster_date_without_snapshot(env):
    env.raw.execute(
        "INSERT INTO jkh_user_inf VALUES(?,?,?,?,?,?,?,?)",
        (
            "later",
            "2026-09-20",
            "004",
            "04",
            "新分行",
            "新支行",
            "later",
            "L1",
        ),
    )
    env.raw.execute("DELETE FROM swe_rm_claw_list_ind_stat")
    response = await request(
        path=f"{BASE_URL}/options",
        query={"end_date": DATE, "kind": "branches"},
    )
    assert response.status_code == 200
    assert response.json()["sync_date"] == DATE
    assert [item["value"] for item in response.json()["items"]] == [
        "001",
        "002",
        "003",
    ]


@pytest.mark.asyncio
async def test_options_fallback_to_latest_roster_date_and_keep_branch_scope(
    env,
):
    env.raw.execute(
        "INSERT INTO jkh_user_inf VALUES(?,?,?,?,?,?,?,?)",
        (
            "later",
            "2026-09-20",
            "004",
            "04",
            "新分行",
            "新支行",
            "later",
            "L1",
        ),
    )
    for requested in ("2026-09-16", "2026-09-30"):
        response = await request(
            path=f"{BASE_URL}/options",
            query={"end_date": requested, "kind": "orgs"},
            headers={"X-Source-Id": "S", "X-Bbk-Id": "004"},
        )
        assert response.status_code == 200
        assert response.json() == {
            "sync_date": "2026-09-20",
            "items": [{"value": "04", "label": "新支行"}],
        }
    forbidden = await request(
        path=f"{BASE_URL}/options",
        query={"end_date": DATE, "kind": "orgs", "first_bbk_id": "004"},
        headers={"X-Source-Id": "S", "X-Bbk-Id": "001"},
    )
    assert forbidden.status_code == 403


@pytest.mark.asyncio
async def test_options_return_empty_when_roster_has_no_dates(env):
    env.raw.execute("DELETE FROM jkh_user_inf")
    response = await request(
        path=f"{BASE_URL}/options",
        query={"end_date": DATE, "kind": "branches"},
    )
    assert response.status_code == 200
    assert response.json() == {"sync_date": None, "items": []}


@pytest.mark.asyncio
async def test_permission_counts_default_missing_and_null_to_zero(
    env, monkeypatch
):
    from monitor.app.services.report import task_type_snapshot as snapshot

    async def missing_counts(*args):
        return {("001", "01"): None}

    monkeypatch.setattr(snapshot, "_fetch_permission_counts", missing_counts)
    response = await request(
        query={
            "end_date": DATE,
            "group_by": "org",
            "task_type": "push_plan",
        }
    )
    assert response.status_code == 200
    assert [
        row["permission_manager_count"] for row in response.json()["items"]
    ] == [0, 0]


@pytest.mark.asyncio
async def test_permission_left_join_keeps_unmatched_sources_but_counts_only_roster(
    env,
):
    from monitor.app.services.report import task_type_snapshot as snapshot

    env.raw.execute(
        "INSERT INTO swe_tenant_init_source VALUES ('unlisted', 'S')"
    )
    env.raw.execute("DELETE FROM jkh_user_inf WHERE user_id = 'alice'")
    counts = await snapshot._fetch_permission_counts(
        env.connection, {"sync_date": DATE}, "S", "manager", {}
    )
    assert counts[("", "", "")] == 0
    assert counts[("002", "01", "bob")] == 1
    report = await service().get_report(params(group_by="manager"), "S", "100")
    assert {
        row.permission_manager_count
        for row in report.items
        if row.user_id == "alice"
    } == {None}
    overall = await service().get_report(params(), "S", "100")
    assert {row.permission_manager_count for row in overall.items} == {1}
    env.raw.execute("DELETE FROM jkh_user_inf")
    empty = await service().get_report(params(), "S", "100")
    assert {row.permission_manager_count for row in empty.items} == {0}


@pytest.mark.asyncio
@pytest.mark.parametrize("group", ["branch", "org", "manager"])
async def test_skill_details_exclude_zero_counts_before_paging_and_export(
    env, group
):
    env.raw.execute("DELETE FROM swe_rm_claw_list_ind_stat")
    dims = {"first_bbk_id": "001", "org_id": "01", "user_id": "alice"}
    rows = [
        snapshot_row(
            f"{group}_skill",
            {**dims, "skill_id": skill},
            "push_plan",
            skill_cnt=count,
        )
        for skill, count in (
            ("a_zero", 0),
            ("b_null", None),
            ("c", 1),
            ("d", 2),
        )
    ]
    rows.append(snapshot_row(group, dims, "push_plan", skill_cnt=0))
    env.raw.executemany(INSERT_SQL, rows)
    query = {
        "end_date": DATE,
        "group_by": group,
        "skill_detail": True,
        "task_type": "push_plan",
    }
    paged = {"page": 1, "page_size": 1} if group == "manager" else {}
    response = await request(query={**query, **paged})
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 2
    assert body["has_more"] is (group == "manager")
    assert [row["skill_id"] for row in body["items"]] == (
        ["c"] if group == "manager" else ["c", "d"]
    )
    if paged:
        second = await request(query={**query, "page": 2, "page_size": 1})
        assert [row["skill_id"] for row in second.json()["items"]] == ["d"]
        assert second.json()["has_more"] is False
    export = await request(path=f"{BASE_URL}/export", query=query)
    assert export.status_code == 200
    sheet = load_workbook(BytesIO(export.content)).active
    assert sheet.max_row == 3
    summary = await request(query={**query, "skill_detail": False})
    assert summary.status_code == 200
    assert summary.json()["total"] == 0
    assert summary.json()["items"] == []
    env.raw.execute("UPDATE swe_rm_claw_list_ind_stat SET skill_cnt = 0")
    empty = await request(query={**query, **paged})
    assert empty.status_code == 200
    assert empty.json()["items"] == []
    assert empty.json()["total"] == 0
    assert empty.json()["has_more"] is False
    assert "no_matching_skills" in empty.json()["warnings"]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "group,detail",
    [
        ("branch", True),
        ("org", True),
        ("manager", True),
        ("manager", False),
    ],
)
async def test_unused_metrics_do_not_query(env, monkeypatch, group, detail):
    from unittest.mock import AsyncMock
    from monitor.app.services.report import task_type_snapshot as snapshot

    permissions = AsyncMock(
        side_effect=AssertionError("unused permission query")
    )
    monkeypatch.setattr(snapshot, "_fetch_permission_counts", permissions)
    queries = []
    env.raw.set_trace_callback(queries.append)
    report = await service().get_report(
        params(group_by=group, skill_detail=detail), "S", "100"
    )
    permissions.assert_not_awaited()
    assert all(row.permission_manager_count is None for row in report.items)
    selects = [
        sql.split(" FROM ")[0]
        for sql in queries
        if "SELECT dw_dat_dt AS prt_dt," in sql and "rpt_combo" in sql
    ]
    assert selects
    if detail:
        assert all("skill_cnt" not in sql for sql in selects)
        assert all(row.skill_count is None for row in report.items)
    if group == "manager":
        assert all("active_manager_cnt" not in sql for sql in selects)


@pytest.mark.asyncio
@pytest.mark.parametrize("detail", [False, True])
@pytest.mark.parametrize("task_type", TASK_TYPES)
async def test_manager_task_counts_and_export(env, detail, task_type):
    env.raw.execute(
        "UPDATE swe_rm_claw_list_ind_stat SET active_job_cnt = 7, "
        "paused_job_cnt = 2 WHERE job_type <> 'ask_plan'"
    )
    query = dict(
        end_date=DATE,
        group_by="manager",
        skill_detail=detail,
        task_type=task_type,
    )
    response = await request(query=query)
    assert response.status_code == 200
    row = response.json()["items"][0]
    expected = [None, None] if task_type == "ask_plan" else [7, 2]
    assert [row["active_task_count"], row["paused_task_count"]] == expected
    assert row["permission_manager_count"] is None
    assert row["active_manager_count"] is None
    exported = await request(path=f"{BASE_URL}/export", query=query)
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content)).active
    headers = [cell.value for cell in sheet[1]]
    assert "有权限客户经理数" not in headers
    assert "活跃客户经理数" not in headers
    assert ("技能数" in headers) is not detail
    start = headers.index("任务类型") + (1 if detail else 2)
    assert headers[start : start + 2] == ["当前活跃任务数", "当前暂停任务数"]
    assert [sheet.cell(2, start + i + 1).value for i in range(2)] == expected


@pytest.mark.asyncio
@pytest.mark.parametrize("group", ["branch", "org", "manager"])
async def test_summary_zero_skills_filtered_before_total_and_export(
    env, group
):
    env.raw.execute("DELETE FROM swe_rm_claw_list_ind_stat")
    rows = [
        snapshot_row(
            group,
            dict(first_bbk_id=f"00{i}", org_id=f"0{i}", user_id=f"user{i}"),
            "push_plan",
            skill_cnt=n,
        )
        for i, n in enumerate([0, None, 1, 2])
    ]
    env.raw.executemany(INSERT_SQL, rows)
    query = dict(end_date=DATE, group_by=group, task_type="push_plan")
    paging = dict(page=1, page_size=1) if group == "manager" else {}
    response = await request(query={**query, **paging})
    assert response.status_code == 200
    body = response.json()
    assert body["total"] == 2
    assert body["has_more"] is (group == "manager")
    assert [row["skill_count"] for row in body["items"]] == (
        [1] if paging else [1, 2]
    )
    exported = await request(path=f"{BASE_URL}/export", query=query)
    sheet = load_workbook(BytesIO(exported.content)).active
    headers = [cell.value for cell in sheet[1]]
    column = headers.index("技能数") + 1
    assert sheet.max_row == 3
    assert [sheet.cell(i, column).value for i in (2, 3)] == [1, 2]


@pytest.mark.asyncio
async def test_org_permission_count_is_scoped_to_branch_and_org(env):
    env.raw.execute(
        "INSERT INTO jkh_user_inf VALUES(?,?,?,?,?,?,?,?)",
        ("dave", DATE, "001", "02", "甲分行", "另一支行", "dave", "L1"),
    )
    env.raw.execute("INSERT INTO swe_tenant_init_source VALUES('dave', 'S')")
    branch = await service().get_report(params(group_by="branch"), "S", "100")
    assert branch.items[0].permission_manager_count == 2
    org = await service().get_report(params(group_by="org"), "S", "100")
    assert {
        (row.first_bbk_id, row.org_id, row.permission_manager_count)
        for row in org.items
    } == {("001", "01", 1), ("002", "01", 1)}
